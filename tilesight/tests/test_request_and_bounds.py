"""Request-level runs (prompt/output, peak memory) and fine-grained bound attribution."""
from dataclasses import replace

from tilesight import HardwareSpec, ModelSpec, RunConfig, _core, run_model
from tilesight.gpuTilingPerfHWModel.model.memory import kv_bytes_per_seq_all
from tilesight.gpuTilingPerfHWModel.interfaceAndRun.request import run_request
from tilesight.gpuTilingPerfHWModel.genResult.table import resource_class

B300, H200 = HardwareSpec.load("b300"), HardwareSpec.load("h200")
KIMI = ModelSpec.load("kimi_k2.hf")

TileConfig = _core.GemmTile
occupancy = _core.occupancy


def lower_gemm(cur_gpu_config, name, M, N, K, **kw):
    kw.setdefault("tile", TileConfig())
    return _core.lower_gemm(cur_gpu_config, name, M, N, K, **kw)


def evaluate(k, hw):
    return _core.evaluate(hw, k)


def test_tpot_grows_with_kv_and_peak_is_prompt_plus_output():
    rc = RunConfig(batch=64, dp=8, prompt_len=4000, output_len=2000, page_size=64, prefill_batch=8)
    r = run_request(KIMI, B300, rc)
    assert r.points[-1].tpot_s > r.points[0].tpot_s
    assert r.points[0].tpot_s <= r.tpot_avg_s <= r.points[-1].tpot_s
    exp = kv_bytes_per_seq_all(KIMI, rc, 6016) * rc.seqs_per_rank / 1e9   # ceil(6000/64)*64
    assert abs(r.peak_kv_GB - exp) < 1e-9
    assert abs(r.e2e_s - (r.ttft_s + r.tpot_avg_s * 2000)) < 1e-9


def test_avg_tpot_matches_dense_sampling():
    rc = RunConfig(batch=64, dp=8, prompt_len=2048, output_len=1024, prefill_batch=8)
    coarse = run_request(KIMI, B300, replace(rc, decode_samples=3))
    fine = run_request(KIMI, B300, replace(rc, decode_samples=9))
    assert abs(coarse.tpot_avg_s - fine.tpot_avg_s) / fine.tpot_avg_s < 0.01


def test_longer_output_needs_more_memory_fewer_requests():
    a = run_request(KIMI, B300, RunConfig(batch=64, dp=8, prompt_len=8192, output_len=1024, prefill_batch=8))
    b = run_request(KIMI, B300, RunConfig(batch=64, dp=8, prompt_len=8192, output_len=32768, prefill_batch=8))
    assert b.peak_kv_GB > a.peak_kv_GB and b.max_requests_per_rank < a.max_requests_per_rank


def test_detail_attributes_tensor():
    rep = run_model(KIMI, B300, RunConfig(phase="decode", batch=256, seq_len=8192, dp=8))
    det = rep.detail_breakdown()
    assert next(iter(det)) == "ddr:load:expert_weight"
    assert any(k.startswith("ddr:load:kv_cache") for k in det)
    # details refine the coarse limiter: totals agree
    assert abs(sum(det.values()) - sum(rep.limiter_breakdown().values())) < 1e-12


def test_naive_softmax_bound_on_scores_tensor():
    rc = RunConfig(phase="prefill", batch=1, seq_len=16384, tp=8, dp=1, attn_impl="naive")
    rep = run_model(ModelSpec.load("llama3_70b.hf"), B300, rc)
    sm = next(o for o in rep.ops if o.op.name.endswith("attn_softmax"))
    assert sm.bottleneck_detail.startswith("ddr:") and ("scores" in sm.bottleneck_detail or "probs" in sm.bottleneck_detail)


def test_register_limited_occupancy_on_hopper():
    # 128x256 fp32 accumulator in registers (no TMEM) -> regs bind on H200, not on B300
    r_h, lim_h = occupancy(H200, 1024, 128 * 256 * 4, 384)
    r_b, lim_b = occupancy(B300, 1024, 128 * 256 * 4, 384)
    assert "regs" in lim_h and "regs" not in lim_b


def test_register_spill_is_charged():
    ks = lower_gemm(H200, "g", 4096, 4096, 4096, tile=TileConfig(bm=128, bn=256, bk=64, stages=2))
    k = ks[0]
    # 128x256 fp32 over 256 consumer threads = 128 acc regs + base: fits in 255, no spill
    assert k.meta["reg_spill_bytes"] == 0 and k.meta["regs_per_thread"] <= 255
    assert "regs" in k.meta["occ_limiter"]


def test_resource_class_parsing():
    assert resource_class("ddr:load:expert_weight") == "memory:DDR/HBM"
    assert resource_class("latency:load:kv_cache(K)(mem-lat)") == "latency<-mem-lat"
    assert resource_class("latency:softmax(sfu)") == "latency<-sfu"
    assert resource_class("smem:mma") == "on-chip:smem"


def test_timeline_matches_engine_round_and_lanes():
    from tilesight.gpuTilingPerfHWModel.genResult.timeline import steady_timeline, timeline_text
    cur_gpu_config = B300
    k = lower_gemm(cur_gpu_config, "g", 4096, 4096, 7168, a_dtype="fp8", b_dtype="fp8", compute_dtype="fp8",
           tile=TileConfig(bm=128, bn=256, bk=64, cluster_m=2))[0]
    tl = steady_timeline(k, cur_gpu_config)
    r = evaluate(k, cur_gpu_config)
    # the timeline's round x iters reproduces the engine's steady time (first wave)
    assert abs(tl["round_s"] - max(tl["resource_bound_s"], tl["latency_bound_s"])) < 1e-15
    assert abs(tl["total_s"] * (k.trace.num_blocks / (cur_gpu_config.sms * tl["resident"])) - r.time_s) / r.time_s < 0.35
    assert tl["limiter"] == "tc" and "tc" in tl["lanes"] and "ddr" in tl["lanes"]
    # every steady item lies inside the drawn window and loads are prefetches
    for it in tl["steady"]:
        assert it["end"] >= it["start"] >= 0
        if it["action"].startswith("load"):
            assert it["iteration"] == it["round"] + k.trace.stages - 1
    txt = timeline_text(tl)
    assert "round" in txt and "tc" in txt


def test_timeline_reflects_the_bottleneck():
    from tilesight.gpuTilingPerfHWModel.genResult.timeline import steady_timeline
    k = _core.lower_attention_decode(B300, "a", B=32, H=64, kv_heads=1, S=8192,
                               d_qk=576, d_v=512, v_in_k=True)[0]
    tl = steady_timeline(k, B300)
    busy = {l: sum(i["lanes"].get(l, 0.0) for i in tl["steady"] if i["round"] == 0) for l in tl["lanes"]}
    top = max(busy, key=busy.get)
    assert top == "ddr"                      # KV streaming dominates decode
    assert busy["ddr"] / tl["round_s"] > 0.7


def test_timeline_cycles_and_trace_text():
    from tilesight.gpuTilingPerfHWModel.genResult.timeline import steady_timeline, timeline_text, trace_text
    cur_gpu_config = B300
    k = lower_gemm(cur_gpu_config, "g", 4096, 4096, 7168, a_dtype="fp8", b_dtype="fp8", compute_dtype="fp8",
           tile=TileConfig(bm=128, bn=256, bk=64, cluster_m=2))[0]
    tl = steady_timeline(k, cur_gpu_config)
    # cycles are seconds x clock, consistently on every item
    assert abs(tl["round_cyc"] - tl["round_s"] * cur_gpu_config.clock_hz) < 1e-9
    for it in tl["steady"]:
        assert abs(it["start_cyc"] - it["start"] * cur_gpu_config.clock_hz) < 1e-6
        for lane, v in it["lanes"].items():
            assert abs(it["lanes_cyc"][lane] - v * cur_gpu_config.clock_hz) < 1e-6
    # wave decomposition is reported (paper §3.4)
    assert tl["full_waves"] * cur_gpu_config.sms * tl["resident"] + tl["tail_blocks"] == k.trace.num_blocks
    # both gantt units render
    assert "cyc" in timeline_text(tl, unit="cyc") and "us" in timeline_text(tl)
    txt = trace_text(tl, "gemm", cur_gpu_config.name)
    for needle in ("start_cyc", "gantt (cycles)", "lane occupancy in one round", "blocks ->"):
        assert needle in txt
    # every steady action appears as an event row
    assert all(a.name in txt for a in k.trace.body)


def test_cycle_csv_one_row_per_cycle_one_column_per_unit():
    import csv as _csv
    import io
    from tilesight.gpuTilingPerfHWModel.genResult.timeline import cycle_csv, machine_timeline, steady_timeline
    cur_gpu_config = B300
    k = lower_gemm(cur_gpu_config, "g", 4096, 4096, 7168, a_dtype="fp8", b_dtype="fp8", compute_dtype="fp8",
           tile=TileConfig(bm=128, bn=256, bk=64, cluster_m=2))[0]
    tl = steady_timeline(k, cur_gpu_config)
    rows = list(_csv.DictReader(io.StringIO(cycle_csv(tl))))
    lanes = tl["lanes"]
    assert set(rows[0]) == {"cycle", "time_ns", "phase", "round"} | set(lanes) | {f"{l}_busy" for l in lanes}
    # consecutive cycles, time matches the clock, busy fractions in [0,1]
    assert [int(r["cycle"]) for r in rows[:50]] == list(range(50))
    assert abs(float(rows[10]["time_ns"]) - 10 / cur_gpu_config.clock_hz * 1e9) < 0.01   # CSV keeps 2 decimals
    for r in rows[:200]:
        for l in lanes:
            b = float(r[f"{l}_busy"])
            assert 0.0 <= b <= 1.0
            assert (b > 0) == bool(r[l])          # a name iff the unit is busy
    # this GEMM is tc-bound: the tc column is occupied by mma most cycles of a round
    tc_busy = sum(float(r["tc_busy"]) for r in rows[:int(tl["round_cyc"])])
    assert tc_busy / tl["round_cyc"] > 0.9
    # full mode covers prologue + every iteration + epilogue
    assert len(cycle_csv(tl, full=True).splitlines()) > len(cycle_csv(tl).splitlines())
    m = machine_timeline(tl)
    assert sum(w["blocks"] for w in m["waves"]) == k.trace.num_blocks
    assert m["waves"][-1]["active_sms"] <= cur_gpu_config.sms


def test_excel_grid_colours_tiles_by_iteration(tmp_path):
    from openpyxl import load_workbook
    from tilesight.gpuTilingPerfHWModel.genResult.excel import write_excel
    from tilesight.gpuTilingPerfHWModel.genResult.timeline import steady_timeline
    cur_gpu_config = B300
    k = lower_gemm(cur_gpu_config, "g", 4096, 4096, 7168, a_dtype="fp8", b_dtype="fp8", compute_dtype="fp8",
           tile=TileConfig(bm=128, bn=256, bk=64, stages=4, cluster_m=2))[0]
    tl = steady_timeline(k, cur_gpu_config)
    assert tl["rounds_shown"] >= k.trace.stages          # load and its consumer both visible
    out = tmp_path / "t.xlsx"
    info = write_excel(tl, str(out), "gemm", cur_gpu_config.name)
    ws = load_workbook(out)["timeline"]
    # row 1 is the unit-section band, row 2 the column names
    assert [c.value for c in ws[2]][:4] == ["cycle", "time_ns", "phase", "round"]
    assert set(tl["lanes"]) <= {c.value for c in ws[2]}
    assert "shader slice · cores" in {c.value for c in ws[1]}
    # a tile keeps its colour when it moves from a memory lane to the tensor core
    cells = {}
    for row in ws.iter_rows(min_row=3, max_row=min(ws.max_row, 6000)):
        for c in row[4:]:
            if c.value and "#" in str(c.value):
                cells.setdefault(str(c.value).split(" ")[0], set()).add(c.fill.start_color.rgb)
    assert all(len(v) == 1 for v in cells.values())          # one colour per tile label
    a_tiles = {k2: v for k2, v in cells.items() if k2.startswith("A#")}
    mmas = {k2: v for k2, v in cells.items() if k2.startswith("MMA#")}
    assert a_tiles and mmas
    for it in (0, 1, 2):
        if f"A#{it}" in a_tiles and f"MMA#{it}" in mmas:
            assert a_tiles[f"A#{it}"] == mmas[f"MMA#{it}"]   # same iteration -> same colour
    assert info["rows"] > 100


def test_extra_onchip_buffer_cuts_hbm_traffic():
    sram = {"memory.sram": {"capacity_MB": 2048, "effective_capacity_MB": 2048,
                            "bandwidth_TBps": 15.0, "latency_ns": 400, "assoc": 16,
                            "per_sm_max_GBps": 180}}
    tile = TileConfig(bm=64, bn=256, bk=64)
    base = lower_gemm(B300, "g", 8, 4096, 7168, batch=48, a_dtype="fp8", b_dtype="fp8",
              compute_dtype="fp8", tile=tile)[0]
    with_buf = lower_gemm(B300.override(sram), "g", 8, 4096, 7168, batch=48, a_dtype="fp8",
                  b_dtype="fp8", compute_dtype="fp8", tile=tile)[0]
    ddr_before = sum(a.work.get("ddr", 0) for a in with_buf.trace.body)
    assert ddr_before < sum(a.work.get("ddr", 0) for a in base.trace.body)
    assert any(a.work.get("sram", 0) > 0 for a in with_buf.trace.body)
    assert evaluate(with_buf, B300.override(sram)).time_s < evaluate(base, B300).time_s


def test_buffer_closed_form_and_search_agree_on_direction():
    from tilesight import ModelSpec, RunConfig
    from tilesight.gpuTilingPerfHWModel.model.dse.buffer import best_alloc, profile_classes, with_buffer
    from tilesight.gpuTilingPerfHWModel.interfaceAndRun.runner import run_model
    m = ModelSpec.load("kimi_k2.hf")
    rc = RunConfig(phase="decode", batch=64, seq_len=4096, dp=8)
    prof = profile_classes(m, B300, rc)
    # value per byte = traffic / footprint, and the fill order follows it
    for c in ("weight", "kv", "act"):
        assert prof["value_per_byte"][c] == prof["traffic_bytes"][c] / prof["footprint_bytes"][c]
    a = best_alloc(prof, 4096)
    order = a["order"]
    assert prof["value_per_byte"][order[0]] >= prof["value_per_byte"][order[-1]]
    assert abs(sum(a["pin"].values()) - 1.0) < 1e-6 or a["unused_bytes"] > 0
    # a buffer never makes it slower, and more capacity never makes it slower
    base = run_model(m, with_buffer(B300, 0), rc).step_time_s
    small = run_model(m, with_buffer(B300, 2048, policy="pin", pin=a["pin"]), rc).step_time_s
    big = run_model(m, with_buffer(B300, 32768, policy="pin", pin=a["pin"]), rc).step_time_s
    assert big <= small <= base + 1e-12


def test_buffer_policies_are_wired():
    from tilesight import ModelSpec, RunConfig
    from tilesight.gpuTilingPerfHWModel.model.dse.buffer import with_buffer
    from tilesight.gpuTilingPerfHWModel.interfaceAndRun.runner import run_model
    m, rc = ModelSpec.load("kimi_k2.hf"), RunConfig(phase="decode", batch=64, seq_len=4096, dp=8)
    plain = run_model(m, with_buffer(B300, 32768, policy="pin", pin={"weight": 1.0}), rc).step_time_s
    tuned = run_model(m, with_buffer(B300, 32768, policy="pin", pin={"weight": 1.0},
                                     bypass_l2=True, prefetch=True), rc).step_time_s
    assert tuned <= plain                      # bypassing L2 / prefetching can only help here
    assert with_buffer(B300, 0).sram is None   # capacity 0 removes the buffer

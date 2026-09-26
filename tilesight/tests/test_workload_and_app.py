"""Workload config (the third input), the wizard's endpoints, PDF / arch diagram / Excel sections."""
import json
import threading
import time
import urllib.request
from http.server import ThreadingHTTPServer

import pytest

from tilesight import HardwareSpec
from tilesight.gpuTilingPerfHWModel.interfaceAndRun.workload import (WORKLOAD_FIELDS, memory_breakdown, run_workload,
                                      run_workload_both_phases, to_run_config, validate_workload)
from tilesight.gpuTilingPerfHWModel.genResult.archdiagram import arch_svg
from tilesight.cli.server import Handler

HW = HardwareSpec.load("b300")
WL = {"name": "demo", "hidden": 7168, "layers": 4,
      "attention": {"type": "mla", "heads": 64},
      "ffn": {"type": "moe", "experts": 48, "topk": 8, "d_ff": 2048},
      "run": {"phase": "decode", "batch": 32, "cur_decoding_seq_len": 4096}}


def test_workload_schema_and_validation():
    assert len(WORKLOAD_FIELDS) >= 30
    assert validate_workload(WL) == []
    bad = validate_workload({"attention": {"type": "nope", "heads": 7, "kv_heads": 2},
                             "dtypes": {"weight": "fp9"}, "typo": 1})
    assert any("unknown field" in p for p in bad)
    assert any("attention.type" in p for p in bad)
    assert any("unknown datatype" in p for p in bad)


def test_workload_is_single_gpu():
    rc = to_run_config(WL)
    assert rc.tp == rc.dp == rc.ep_size == 1 and not rc.include_lm_head
    rep = run_workload(WL, HW)
    assert rep.step_time_s > 0 and rep.ops
    assert not any(o.op.kind in ("allreduce", "a2a") for o in rep.ops)   # no multi-GPU traffic


@pytest.mark.parametrize("attn,ffn", [("mla", "moe"), ("gqa", "mlp"), ("mha", "none")])
def test_every_attention_and_ffn_combination_runs(attn, ffn):
    cfg = dict(WL, attention={"type": attn, "heads": 16, "kv_heads": 4, "head_dim": 128},
               ffn={"type": ffn, "d_ff": 4096, "experts": 8, "topk": 2})
    assert validate_workload(cfg) == []
    assert run_workload(cfg, HW).step_time_s > 0


def test_arch_diagram_comes_from_the_config():
    svg = arch_svg(HW)
    assert svg.startswith("<svg") and "http" not in svg
    assert f"{HW.sms} SMs" in svg and "HBM" in svg and "L2 partition 0" in svg
    # a different config draws a different picture
    assert arch_svg(HardwareSpec.load("mi355x")) != svg


def test_pdf_report(tmp_path):
    from tilesight import _core
    lower_gemm = _core.lower_gemm
    TileConfig = _core.GemmTile
    from tilesight.gpuTilingPerfHWModel.genResult.pdfreport import unit_groups, write_pdf
    ks = lower_gemm(HW, "gemm", 4096, 4096, 7168, a_dtype="fp8", b_dtype="fp8",
                    compute_dtype="fp8", tile=TileConfig(128, 256, 64))
    out = tmp_path / "r.pdf"
    write_pdf(str(out), HW, ks, "test", WL, run_workload(WL, HW))
    data = out.read_bytes()
    assert data[:5] == b"%PDF-" and len(data) > 3000
    g = unit_groups(["tc", "cuda", "smem", "l1", "sram", "path:tma", "l2", "ddr"])
    # the three blocks of the model: shader slice / gmem / memory
    assert g["shader slice · cores"] == ["tc", "cuda"]
    assert g["shader slice · L1/scratchpad"] == ["smem", "l1"]
    assert g["on-chip buffer"] == ["sram"] and g["shader slice · load paths"] == ["path:tma"]
    assert g["memory · L2 ports"] == ["l2"] and g["memory · HBM"] == ["ddr"]


def test_excel_columns_are_grouped_by_unit(tmp_path):
    from openpyxl import load_workbook
    from tilesight import _core
    lower_gemm = _core.lower_gemm
    TileConfig = _core.GemmTile
    from tilesight.gpuTilingPerfHWModel.genResult.excel import write_excel
    from tilesight.gpuTilingPerfHWModel.genResult.timeline import steady_timeline
    k = lower_gemm(HW, "g", 4096, 4096, 7168, a_dtype="fp8", b_dtype="fp8",
                   compute_dtype="fp8", tile=TileConfig(128, 256, 64))[0]
    out = tmp_path / "t.xlsx"
    write_excel(steady_timeline(k, HW), str(out), "g", HW.name)
    ws = load_workbook(out)["timeline"]
    assert [c.value for c in ws[1]][:5] == [None, None, None, None, "memory · HBM"]
    assert [c.value for c in ws[2]][:5] == ["cycle", "time_ns", "phase", "round", "ddr"]
    assert ws.cell(row=3, column=1).value == 0
    assert ws.freeze_panes == "E3"


@pytest.fixture(scope="module")
def url():
    srv = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_address[1]}"
    srv.shutdown()


def _job(url, mode, cfg):
    req = urllib.request.Request(url + "/api/jobs", method="POST",
                                 data=json.dumps({"mode": mode, "config": cfg}).encode(),
                                 headers={"Content-Type": "application/json"})
    jid = json.load(urllib.request.urlopen(req))["job_id"]
    for _ in range(900):
        j = json.load(urllib.request.urlopen(f"{url}/api/jobs/{jid}"))
        if j["state"] != "running":
            return jid, j
        time.sleep(0.05)
    raise AssertionError("timeout")


def test_wizard_endpoints(url):
    page = urllib.request.urlopen(url + "/").read().decode()
    assert "Step 1" in page and "Step 5" in page and "http://" not in page.split("<script>")[0]
    sc = json.load(urllib.request.urlopen(url + "/api/schema"))
    assert len(sc["hardware"]) > 90 and len(sc["workload"]) >= 30
    assert urllib.request.urlopen(url + "/api/gpu_tiling_perf_hw_model_yaml?gpuTilingPerfHWModel=b300").read().startswith(b"#")
    assert urllib.request.urlopen(url + "/api/arch.svg?gpuTilingPerfHWModel=b300").read().startswith(b"<svg")
    req = urllib.request.Request(url + "/api/validate_gpu_tiling_perf_hw_model", method="POST",
                                 data=json.dumps({"yaml": "name: x\nsms: 4\n"}).encode(),
                                 headers={"Content-Type": "application/json"})
    assert any("missing required" in p for p in json.load(urllib.request.urlopen(req))["problems"])


def test_workload_job_reports_progress_and_downloads(url):
    jid, j = _job(url, "workload", {"gpuTilingPerfHWModel": "b300", "workload": WL})
    assert j["state"] == "done"
    r = j["result"]
    assert r["arch_svg"].startswith("<svg") and r["timeline"] and r["trace_kernel"]
    assert any(":" in line and "→" in line for line in j["log"])      # file:function shown live
    for ep, ctype in (("xlsx", "spreadsheet"), ("csv", "text/csv"), ("pdf", "application/pdf")):
        resp = urllib.request.urlopen(f"{url}/api/{ep}?job={jid}")
        assert ctype in resp.headers["Content-Type"] and len(resp.read()) > 1000


def test_custom_gpu_yaml_is_accepted_and_validated(url):
    y = urllib.request.urlopen(url + "/api/gpu_tiling_perf_hw_model_yaml?gpuTilingPerfHWModel=h200").read().decode()
    _, ok = _job(url, "workload", {"gpuTilingPerfHWModelYaml": y.replace("sms: 132", "sms: 99"), "workload": WL})
    assert ok["state"] == "done" and ok["result"]["gpu_name"] == "H200"
    _, bad = _job(url, "workload", {"gpuTilingPerfHWModelYaml": "name: broken\nsms: 8\n", "workload": WL})
    assert bad["state"] == "error" and "hardware config" in bad["error"]


def test_slice_config_derives_and_translates():
    import yaml
    from tilesight.gpuTilingPerfHWModel.interfaceAndRun.slice_config import (derive, to_hardware_spec, validate_slice_config)
    cfg = yaml.safe_load(open("examples/slice_gpu.yaml"))
    assert validate_slice_config(cfg) == []
    d = derive(cfg)
    # 2*64*64*32 flops per tile / 16 cycles * 2 GHz = 32.768 TFLOP/s per tensor core
    assert abs(d["tflops_per_tensor_core"] - 2 * 64 * 64 * 32 / 16 * 2e9 / 1e12) < 1e-6
    assert d["shader_cores"] == 8 * 20 and d["tensor_cores"] == 8 * 20 * 4
    assert abs(d["pflops_total"] - d["tflops_per_tensor_core"] * d["tensor_cores"] / 1000) < 1e-6
    assert d["hbm_capacity_GB"] == 8 * 24 and abs(d["hbm_TBps"] - 8.0) < 1e-9
    cur_gpu_config = to_hardware_spec(cfg)
    assert cur_gpu_config.validate() == [] and cur_gpu_config.sms == 160
    assert cur_gpu_config.get("memory.l2.partitions") == 8          # one per memory slice
    assert cur_gpu_config.get("compute.attention_tile_m") == 64
    # L2 = 0 rewires the DMA straight to the buffer
    z = to_hardware_spec({**cfg, "memory_slice": {**cfg["memory_slice"], "l2_MB": 0}})
    assert z.dma_destination == "bypass"
    # the buffer contents policy maps to pinned classes
    for contents, cls in (("ab", "act"), ("abc_kv", "kv"), ("abc_kv_moe", "weight")):
        h = to_hardware_spec({**cfg, "onchip_buffer": {**cfg["onchip_buffer"], "contents": contents}})
        assert cls in h.get("memory.sram.pin")
    assert to_hardware_spec({**cfg, "onchip_buffer": {**cfg["onchip_buffer"],
                                                      "contents": "none"}}).sram is None


def test_kimi_k3_preset_memory_matches_the_simulation():
    import yaml
    from tilesight.gpuTilingPerfHWModel.interfaceAndRun.workload import memory_breakdown
    cfg = yaml.safe_load(open("modelPresets/kimi_k3_10L.yaml"))
    # max_seq_len sizes the KV cache independently of the length an actual run samples at
    # (that's the point — see workload.py's memory_breakdown()); run this comparison with
    # cur_decoding_seq_len one below max_seq_len (validate_workload() requires max_seq_len
    # strictly greater than prefill_seq_len + cur_decoding_seq_len) so it stays a check of the
    # KV-bytes formula, not of the two knobs' (deliberately) different defaults — one token out
    # of 16384 is well inside the 5% tolerance below.
    cfg["run"]["prefill_seq_len"] = 0
    cfg["run"]["cur_decoding_seq_len"] = 16383
    cfg["run"]["max_seq_len"] = 16384
    assert validate_workload(cfg) == []
    m = memory_breakdown(cfg)
    assert m["layers"] == 10 and "16/112" in m["sparsity"]
    assert m["kv_bytes_per_token_per_layer"] == (512 + 64) * 2          # latent + rope, bf16
    rep = run_workload(cfg, HW)
    assert abs(m["weights_GB"] - rep.memory.weights_GB) / rep.memory.weights_GB < 0.05
    assert abs(m["kv_cache_GB"] - rep.memory.kv_cache_GB) / rep.memory.kv_cache_GB < 0.05


def test_max_seq_len_must_exceed_prefill_plus_decoding():
    ok = {**WL, "run": {**WL["run"], "prefill_seq_len": 4096, "cur_decoding_seq_len": 8192,
                        "max_seq_len": 16384}}
    assert validate_workload(ok) == []
    too_small = {**WL, "run": {**WL["run"], "prefill_seq_len": 4096, "cur_decoding_seq_len": 8192,
                               "max_seq_len": 12288}}          # exactly equal, not greater
    problems = validate_workload(too_small)
    assert len(problems) == 1 and "max_seq_len" in problems[0]


def test_run_workload_both_phases_runs_prefill_and_decode_at_the_given_lengths():
    cfg = {**WL, "run": {**WL["run"], "prefill_seq_len": 2048, "cur_decoding_seq_len": 6000,
                         "max_seq_len": 16384}}
    reps = run_workload_both_phases(cfg, HW)
    assert reps["prefill"].rc.phase == "prefill" and reps["prefill"].rc.seq_len == 2048
    assert reps["decode"].rc.phase == "decode" and reps["decode"].rc.seq_len == 6000
    # decode-at-a-length is a single token step: far cheaper than prefilling the whole prompt
    assert reps["decode"].step_time_s < reps["prefill"].step_time_s


def test_run_workload_both_phases_rejects_an_invalid_config():
    bad = {**WL, "run": {**WL["run"], "prefill_seq_len": 8192, "cur_decoding_seq_len": 8192,
                         "max_seq_len": 8192}}
    with pytest.raises(ValueError, match="max_seq_len"):
        run_workload_both_phases(bad, HW)


def test_memory_breakdown_reports_on_chip_buffer_kv_fit_when_given_hw():
    cfg = {**WL, "run": {**WL["run"], "prefill_seq_len": 4096, "cur_decoding_seq_len": 8192,
                         "max_seq_len": 16384}}
    plain = memory_breakdown(cfg)
    assert "kv_fits_on_chip_buffer" not in plain              # no cur_gpu_config -> no opinion
    with_hw = memory_breakdown(cfg, HW)
    assert "kv_fits_on_chip_buffer" in with_hw and "on_chip_buffer_kv_capacity_GB" in with_hw
    # an on-chip buffer big enough for this workload's whole KV cache -> the assumption holds
    big_buffer = HW.override({"memory.sram.capacity_MB": with_hw["kv_cache_GB"] * 1024 * 2,
                              "memory.sram.alloc": {"kv": 1.0}})
    assert memory_breakdown(cfg, big_buffer)["kv_fits_on_chip_buffer"]


def test_derive_endpoint_and_slice_job(url):
    import yaml
    cfg = yaml.safe_load(open("examples/slice_gpu.yaml"))
    wl = yaml.safe_load(open("modelPresets/kimi_k3_10L.yaml"))
    req = urllib.request.Request(url + "/api/derive", method="POST",
                                 data=json.dumps({"slice_cfg": cfg, "workload": wl}).encode(),
                                 headers={"Content-Type": "application/json"})
    d = json.load(urllib.request.urlopen(req))
    assert d["problems"] == [] and d["derived"]["pflops_total"] > 0
    assert d["arch_svg"].startswith("<svg") and d["memory"]["weights_GB"] > 0
    assert len(json.load(urllib.request.urlopen(url + "/api/slice_schema"))["fields"]) >= 40
    _, j = _job(url, "workload", {"slice_cfg": cfg, "workload": wl})
    assert j["state"] == "done" and j["result"]["derived"]["pflops_total"] > 0
    assert j["result"]["workload_memory"]["total_GB"] > 0


def test_model_is_organised_in_three_blocks():
    from tilesight.gpuTilingPerfHWModel.interfaceAndRun.hardware_spec import lane_domain
    from tilesight.gpuTilingPerfHWModel.genResult.table import domain_of
    lanes = {l.name: l for l in HW.lanes()}
    assert lanes["tc"].domain == "shader_slice" and lanes["tc"].scope == "per_core"
    assert lanes["l1"].domain == "shader_slice" and lanes["l1"].scope == "per_slice"
    assert lanes["path:tma"].domain == "shader_slice" and lanes["l2"].domain == "memory"
    assert lane_domain("sram") == "onchip_buffer" and lane_domain("ddr") == "memory"
    # L1 is a slice resource: its modelled total is slices x per-slice bandwidth
    slices = HW.sms // HW.get("memory.l1.cluster_size")
    per_slice = HW.get("memory.l1.bytes_per_clk") * HW.clock_hz * HW.eff("smem")
    assert abs(lanes["l1"].total_rate - per_slice * slices) < 1
    # limiters roll up into the same three blocks
    assert domain_of("ddr:load:weight") == "memory"
    assert domain_of("tc:mma") == "shader_slice"
    assert domain_of("path:tma:load:act") == "shader_slice"
    rep = run_workload(WL, HW)
    blocks = rep.by_domain()
    assert set(blocks) <= {"shader_slice", "onchip_buffer", "memory", "runtime"}
    assert abs(sum(blocks.values()) - sum(rep.detail_breakdown().values())) < 1e-12


def test_flash_attention_is_a_choice_with_consequences():
    from tilesight.gpuTilingPerfHWModel.interfaceAndRun.workload import compare_attention_impl
    cfg = dict(WL, run={**WL["run"], "phase": "prefill", "batch": 1, "prefill_seq_len": 8192})
    c = compare_attention_impl(cfg, HW)
    assert c["flash"]["step_ms"] < c["naive"]["step_ms"]
    assert c["speedup_of_flash"] > 1.2
    assert c["extra_activation_GB_without_flash"] > 1.0        # S and P must live in HBM
    assert c["naive"]["attention_ms"] > c["flash"]["attention_ms"] * 2
    # and the blocks shift: flash leans on the shader slice, naive on memory
    assert c["flash"]["by_block"]["shader_slice"] > c["naive"]["by_block"]["shader_slice"]


def test_buffer_sits_between_the_slices_and_l2_behind_a_switch():
    from tilesight import _core
    from tilesight.gpuTilingPerfHWModel.model.dse.buffer import with_buffer
    lower_gemm = _core.lower_gemm
    TileConfig = _core.GemmTile
    tile = TileConfig(bm=64, bn=128, bk=64)
    plain = lower_gemm(HW, "g", 8, 16384, 7168, a_dtype="fp8", b_dtype="fp8",
                       compute_dtype="fp8", tile=tile)[0]
    buf_hw = with_buffer(HW, 64 * 1024, policy="pin", pin={"weight": 1.0}, prefetch=True)
    buf = lower_gemm(buf_hw, "g", 8, 16384, 7168, a_dtype="fp8", b_dtype="fp8",
                     compute_dtype="fp8", tile=tile)[0]
    w = lambda k, lane: sum(a.work.get(lane, 0) for a in k.trace.body)     # noqa: E731
    # a tile answered by the buffer crosses the switch and touches neither L2 nor HBM
    assert w(buf, "switch") > 0 and w(buf, "sram") == w(buf, "switch")
    assert w(buf, "ddr") < w(plain, "ddr") and w(buf, "l2") < w(plain, "l2")
    assert w(plain, "switch") == 0
    lanes = {l.name: l for l in buf_hw.lanes()}
    assert lanes["switch"].domain == "onchip_buffer" and lanes["sram"].domain == "onchip_buffer"
    # the switch has an aggregate bandwidth and a per-slice share
    assert lanes["switch"].per_sm_cap < lanes["switch"].total_rate
    # and it can be the bottleneck on its own
    narrow = buf_hw.override({"memory.sram.switch_TBps": 0.5})
    wide = buf_hw.override({"memory.sram.switch_TBps": 40.0})
    tn = _core.evaluate(narrow, lower_gemm(narrow, "g", 8, 16384, 7168, a_dtype="fp8", b_dtype="fp8",
                                     compute_dtype="fp8", tile=tile)[0]).time_s
    tw = _core.evaluate(wide, lower_gemm(wide, "g", 8, 16384, 7168, a_dtype="fp8", b_dtype="fp8",
                                     compute_dtype="fp8", tile=tile)[0]).time_s
    assert tn > tw


def test_no_l2_means_ports_straight_to_hbm_not_a_throttled_lane():
    from tilesight.gpuTilingPerfHWModel.model.dse.buffer import with_buffer
    no_l2 = HW.override({"memory.l2.capacity_MB": 0})
    lanes = {l.name: l for l in no_l2.lanes()}
    ddr_rate = no_l2.get("memory.ddr.bandwidth_TBps") * 1e12 * no_l2.eff("ddr")
    assert lanes["l2"].total_rate >= ddr_rate            # the ports are not a false bottleneck
    a = run_workload(WL, no_l2).step_time_s
    b = run_workload(WL, HW).step_time_s
    assert b < a < 3 * b                                 # losing L2 hurts, but does not explode
    # with the weights in the buffer, skipping L2 entirely is fine
    c = run_workload(WL, with_buffer(no_l2, 64 * 1024, policy="pin", pin={"weight": 1.0},
                                     prefetch=True)).step_time_s
    assert c < a
